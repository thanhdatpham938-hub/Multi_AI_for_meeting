"""Reviews (accept / edit / reject, speaker renames) and JSON / Excel export."""

from __future__ import annotations

import html
import json
import re
import unicodedata
from io import BytesIO
from pathlib import Path
from typing import Any, NamedTuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Item, ItemType, RunResult, Utterance
from .signals import TAG_LABELS  # single source of tag display names

REVIEWS_FILE = "reviews.json"
EXPORTED_REVIEWS = {"accepted", "edited"}
EDITABLE_FIELDS = ("content", "status", "owner", "deadline", "condition", "pending_on", "follow_up_owner")
# Other review keys ("note", "answers", "resolved") are reviewer bookkeeping: apply_reviews ignores them
# and they stay in reviews.json only.

TYPE_LABELS = {
    ItemType.decision: "Quyết định",
    ItemType.action: "Việc cần làm",
    ItemType.blocker: "Vướng mắc",
    ItemType.request: "Yêu cầu",
    ItemType.open_issue: "Vấn đề còn mở",
    ItemType.knowledge: "Tri thức mới",
}
STATUS_LABELS = {
    "proposed": "Đề xuất",
    "superseded": "Bị thay thế",
    "confirmed": "Đã chốt",
    "conditional": "Có điều kiện",
    "deferred": "Hoãn",
    "open": "Chưa chốt",
}
FLAG_LABELS = {
    "no_closing_signal": "Không có tín hiệu chốt",
    "authority_unclear": "Chưa rõ thẩm quyền",
    "number_check": "Kiểm tra con số",
    "missing_owner": "Thiếu người phụ trách",
    "verify_failed": "Kiểm chứng không đạt",
    "low_confidence": "Độ tin cậy thấp",
    "tracker_failed": "Lỗi theo dõi quyết định",
    "conflicting_decision": "Quyết định mâu thuẫn",
    "unresolved_reference": "Tham chiếu chưa rõ",
    "needs_recheck": "Cần kiểm tra lại",
    "quote_mismatch": "Trích dẫn không khớp",
    "signal_mismatch": "Lệch tín hiệu code/AI",
    "possible_action": "Có thể là việc cần làm",
    "possible_closing": "Có thể đã chốt",
    "possible_superseded": "Có thể đã bị thay thế",
    "authority_unconfirmed": "Người chốt không có quyền, chưa ai xác nhận",
    "number_not_in_evidence": "Số không có trong câu gốc",
}
COMMITMENT_LABELS = {
    "assigned": "Được giao việc",
    "accepted": "Đã nhận việc",
    "volunteered_conditional": "Xung phong (có điều kiện)",
    "already_done": "Đã làm xong",
    "hypothetical": "Giả định",
}
SCOPE_LABELS = {"ephemeral": "Tạm thời", "project": "Theo dự án", "reusable": "Dùng lại nhiều lần"}
KNOWLEDGE_SCOPES_KEPT = ("project", "reusable")
KNOWLEDGE_SHEET = "Kho tri thức"
REVIEW_LABELS = {"pending": "Chờ duyệt", "accepted": "Chấp nhận", "edited": "Đã sửa", "rejected": "Loại"}


# ---------------------------------------------------------------- reviews


def empty_reviews() -> dict[str, Any]:
    return {"items": {}, "speaker_names": {}}


def load_reviews(run_dir: str | Path) -> dict[str, Any]:
    path = Path(run_dir) / REVIEWS_FILE
    if not path.exists():
        return empty_reviews()
    data = json.loads(path.read_text(encoding="utf-8"))
    return {"items": data.get("items", {}), "speaker_names": data.get("speaker_names", {})}


def save_reviews(run_dir: str | Path, reviews: dict[str, Any]) -> Path:
    path = Path(run_dir) / REVIEWS_FILE
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(reviews, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def rename_speakers(text: str | None, names: dict[str, str]) -> str | None:
    if not text or not names:
        return text
    for label, real in names.items():
        if real.strip():
            text = re.sub(rf"\b{re.escape(label)}\b", real.strip(), text)
    return text


def apply_reviews(result: RunResult, reviews: dict[str, Any]) -> RunResult:
    """Return a copy of `result` with review decisions, edits and speaker names applied."""
    names = {k: v for k, v in reviews.get("speaker_names", {}).items() if v and v.strip()}
    out = result.model_copy(deep=True)
    for u in out.utterances:
        u.speaker = rename_speakers(u.speaker, names) or u.speaker
    items: list[Item] = []
    for item in out.items:
        rev = reviews.get("items", {}).get(item.item_id, {})
        updates: dict[str, Any] = {}
        for f in EDITABLE_FIELDS:
            if f in rev:
                updates[f] = rev[f] or None
        if "review" in rev:
            updates["review"] = rev["review"]
        item = item.model_copy(update=updates)
        item = Item.model_validate(item.model_dump())  # re-validate edited status values
        if item.owner and "missing_owner" in item.flags:  # the reviewer supplied the owner
            item.flags = [f for f in item.flags if f != "missing_owner"]
        item.owner = rename_speakers(item.owner, names)
        item.content = rename_speakers(item.content, names) or item.content
        items.append(item)
    out.items = items
    return out


# ---------------------------------------------------------------- export


def to_json(result: RunResult) -> str:
    return json.dumps(result.model_dump(mode="json"), ensure_ascii=False, indent=2)


def _evidence_text(item: Item, result: RunResult) -> str:
    by_id = {u.id: u for u in result.utterances}
    return "\n".join(
        f"{e} [{by_id[e].time}] {by_id[e].speaker}: {by_id[e].text}" for e in item.evidence if e in by_id
    )


def _history_text(item: Item) -> str:
    return "\n".join(
        f"{STATUS_LABELS.get(h.status.value, h.status.value)}: {h.content or ''}"
        f"{f' ({h.reason})' if h.reason else ''} [{', '.join(h.evidence)}]"
        for h in item.history
    )


COLUMNS = [
    ("ID", 7), ("Chủ đề", 22), ("Nội dung", 48), ("Trạng thái", 14), ("Điều kiện", 24),
    ("Người phụ trách", 16), ("Hạn chót", 14), ("Độ tin cậy", 11), ("Cờ", 28),
    ("Trích dẫn", 60), ("Lịch sử", 50), ("Duyệt", 12),
]


_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def append_row(ws: Any, row: list[Any]) -> None:
    """Append a row; text starting with = + - @ is stored as a string, never as a formula."""
    ws.append(row)
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, str) and cell.value.startswith(_FORMULA_PREFIXES):
            cell.data_type = "s"


def to_excel(result: RunResult, only_reviewed: bool = True) -> bytes:
    """One sheet per item type. By default only accepted/edited items are exported."""
    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    info = wb.create_sheet("Thông tin")
    meta = result.meta
    rows = [
        ("Cuộc họp", (meta.title if meta else None) or result.source_file),
        ("Ngày", meta.date if meta else None),
        ("Thời gian", f"{meta.start_time or ''} – {meta.end_time or ''}" if meta and (meta.start_time or meta.end_time) else None),
        ("Dự án", meta.project if meta else None),
        ("Chương trình", "; ".join(meta.agenda) if meta else None),
        ("Người tham dự", ", ".join(p.name for p in meta.participants) if meta else None),
        ("Run ID", result.run_id),
        ("Model", f"{result.model_light} / {result.model_strong}"),
    ]
    for r in rows:
        append_row(info, [r[0], r[1] or ""])
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 70

    for item_type, label in TYPE_LABELS.items():
        items = [
            i for i in result.items
            if i.type == item_type and (not only_reviewed or i.review in EXPORTED_REVIEWS)
        ]
        if not items:
            continue
        ws = wb.create_sheet(label)
        ws.append([c for c, _ in COLUMNS])
        for idx, (_, width) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=1, column=idx)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            ws.column_dimensions[get_column_letter(idx)].width = width
        for i in items:
            append_row(ws, [
                i.item_id, i.topic, i.content,
                STATUS_LABELS.get(i.status.value, i.status.value) if i.status else "",
                i.condition or "", i.owner or "", i.deadline or "", round(i.confidence, 2),
                ", ".join(FLAG_LABELS.get(f, f) for f in i.flags),
                _evidence_text(i, result), _history_text(i), REVIEW_LABELS.get(i.review, i.review),
            ])
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
    _add_knowledge_sheet(wb, result, only_reviewed, header_fill)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_run(run_dir: str | Path, fmt: str = "xlsx", only_reviewed: bool = True) -> Path:
    from .pipeline import load_result

    run_dir = Path(run_dir)
    result = apply_reviews(load_result(run_dir), load_reviews(run_dir))
    if fmt == "json":
        path = run_dir / "export.json"
        path.write_text(to_json(result), encoding="utf-8")
    elif fmt == "knowledge":
        path = run_dir / "knowledge.json"
        path.write_text(knowledge_to_json(result, only_reviewed=only_reviewed), encoding="utf-8")
    elif fmt == "xlsx":
        path = run_dir / "export.xlsx"
        path.write_bytes(to_excel(result, only_reviewed=only_reviewed))
    else:
        raise ValueError(f"Định dạng không hỗ trợ: {fmt}")
    return path


# ---------------------------------------------------------------- safe text for st.markdown

_MD_ENTITIES = {"\\": "&#92;", "`": "&#96;", "*": "&#42;", "_": "&#95;", "[": "&#91;", "]": "&#93;", "$": "&#36;",
                "~": "&#126;"}


def md_safe(text: Any) -> str:
    """html.escape plus numeric entities for markdown / LaTeX syntax (dollar, asterisk, underscore, brackets,
    tilde, backslash, backtick).

    Entities (not backslashes) so the result is also literal inside raw HTML blocks passed to st.markdown.
    """
    return "".join(_MD_ENTITIES.get(ch, ch) for ch in html.escape(str(text)))


# ---------------------------------------------------------------- evidence display (pure helpers)


def _fold_with_map(text: str) -> tuple[str, list[int]]:
    """Lowercase + collapse whitespace like provenance.normalize_text (without edge stripping).

    Returns the folded string and, for every folded char, the index in `text` it came from.
    """
    folded: list[str] = []
    index: list[int] = []
    prev_space = True  # drops leading whitespace
    for pos, ch in enumerate(text):
        if ch.isspace():
            if not prev_space:
                folded.append(" ")
                index.append(pos)
            prev_space = True
            continue
        prev_space = False
        for low in ch.lower():
            folded.append(low)
            index.append(pos)
    return "".join(folded), index


def locate_quote(text: str, quote: str) -> tuple[int, int] | None:
    """Span (start, end) of `quote` inside NFC-normalized `text`, matched after normalization; None if absent."""
    from .provenance import normalize_text  # local import keeps export light

    q = normalize_text(quote)
    if not q:
        return None
    folded, index = _fold_with_map(unicodedata.normalize("NFC", text))
    start = folded.find(q)
    if start < 0:
        return None
    return index[start], index[start + len(q) - 1] + 1


def highlight_quote_html(text: str, quote: str | None) -> tuple[str, bool | None]:
    """HTML-escaped `text` with the quote in <b>. Flag: None = no quote given, True = found, False = not found."""
    text = unicodedata.normalize("NFC", text)
    if not quote or not quote.strip():
        return md_safe(text), None
    span = locate_quote(text, quote)
    if span is None:
        return md_safe(text), False
    a, b = span
    return f"{md_safe(text[:a])}<b>{md_safe(text[a:b])}</b>{md_safe(text[b:])}", True


def utterance_context(utterances: list[Utterance], uid: str, n: int = 2) -> tuple[list[Utterance], list[Utterance]]:
    """Up to `n` utterances before and after `uid`; ([], []) when the id is unknown."""
    pos = next((k for k, u in enumerate(utterances) if u.id == uid), None)
    if pos is None:
        return [], []
    return utterances[max(0, pos - n):pos], utterances[pos + 1:pos + 1 + n]


def related_summaries(item: Item, result: RunResult, max_len: int = 80) -> list[tuple[str, str, bool]]:
    """(item_id, short content, exists) for each related item; unknown ids are kept with exists=False."""
    by_id = {i.item_id: i for i in result.items}
    out: list[tuple[str, str, bool]] = []
    for rid in item.related_items:
        other = by_id.get(rid)
        if other is None:
            out.append((rid, "", False))
            continue
        content = other.content if len(other.content) <= max_len else other.content[: max_len - 1].rstrip() + "…"
        out.append((rid, content, True))
    return out


# ---------------------------------------------------------------- clarification questions (pure helpers)

CLARIFY_FLAGS = ("missing_owner", "unresolved_reference", "conflicting_decision", "needs_recheck", "quote_mismatch",
                 "signal_mismatch", "possible_action", "possible_closing", "possible_superseded",
                 "authority_unconfirmed", "number_not_in_evidence")
# High-level flags: shown as prominent badges; every other flag is a low-level note (collapsed in the table).
HIGH_LEVEL_FLAGS = ("possible_superseded", "authority_unconfirmed", "number_not_in_evidence", "signal_mismatch",
                    "verify_failed", "conflicting_decision", "tracker_failed")


def split_flags(flags: list[str]) -> tuple[list[str], list[str]]:
    """(high-level flags, low-level flags), each keeping the original order."""
    return ([f for f in flags if f in HIGH_LEVEL_FLAGS], [f for f in flags if f not in HIGH_LEVEL_FLAGS])


def signal_notes(item: Item) -> list[dict[str, Any]]:
    """Well-formed notes from item.signal_check (tolerates empty / malformed data)."""
    check = item.signal_check if isinstance(item.signal_check, dict) else {}
    notes = check.get("notes")
    return [n for n in notes if isinstance(n, dict)] if isinstance(notes, list) else []
NO_FOLLOW_UP = "no_follow_up"  # pseudo-reason: deferred / conditional item nobody follows up
_NUMBER_RE = re.compile(
    r"\d[\d.,/]*\s?(?:triệu|tỷ|nghìn|ngàn|%|đồng|vnđ|vnd|usd|k|phút|giờ|ngày|tuần|tháng)?(?!\w)", re.IGNORECASE
)


class Clarification(NamedTuple):
    item_id: str
    reason: str  # one of CLARIFY_FLAGS or NO_FOLLOW_UP
    question: str
    answer_field: str  # "owner" | "follow_up_owner" | "note"
    answered: str | None


def _short(text: str, n: int = 90) -> str:
    text = " ".join(text.split())
    return text if len(text) <= n else text[: n - 1].rstrip() + "…"


def _first_quote(item: Item, result: RunResult) -> str:
    for quote in item.evidence_quotes.values():
        if quote.strip():
            return _short(quote)
    by_id = {u.id: u for u in result.utterances}
    for uid in item.evidence:
        if uid in by_id:
            return _short(by_id[uid].text)
    return _short(item.content)


def _number_in(text: str) -> str | None:
    match = _NUMBER_RE.search(text)
    return match.group(0).strip() if match else None


def _other_decision_text(item: Item, result: RunResult) -> str:
    """Content of the decision that conflicts with `item`: from conflict_pairs, else related_items, else generic."""
    from .conflicts import conflict_pairs

    for a, b in conflict_pairs(result.items):
        if a.item_id == item.item_id:
            return _short(b.content)
        if b.item_id == item.item_id:
            return _short(a.content)
    by_id = {i.item_id: i for i in result.items}
    for rid in item.related_items:
        other = by_id.get(rid)
        if other is not None and other.type == ItemType.decision and other.item_id != item.item_id:
            return _short(other.content)
    return "một quyết định khác trong cuộc họp"


def _utt_text(result: RunResult, uid: Any) -> tuple[str, str] | None:
    """(text, time) of utterance `uid`, or None when unknown."""
    for u in result.utterances:
        if u.id == uid:
            return u.text, u.time
    return None


def _closing_uid(item: Item) -> str | None:
    check = item.signal_check if isinstance(item.signal_check, dict) else {}
    uid = check.get("closing_utterance")
    return uid if isinstance(uid, str) and uid else (item.evidence[-1] if item.evidence else None)


def _note_utterance(item: Item, rules: tuple[str, ...]) -> str | None:
    for n in signal_notes(item):
        if n.get("rule") in rules and isinstance(n.get("utterance"), str) and n["utterance"]:
            return n["utterance"]
    return None


def _signal_clarification(item: Item, reason: str, result: RunResult, content: str) -> str | None:
    """Question for the signal-layer flags; falls back to a generic question when data is missing."""
    check = item.signal_check if isinstance(item.signal_check, dict) else {}
    closing = _utt_text(result, _closing_uid(item))
    if reason == "signal_mismatch":
        tags = check.get("tags") if isinstance(check.get("tags"), dict) else {}
        labels = [TAG_LABELS.get(code, str(phrase)) for code, phrase in tags.items()]
        if closing and labels:
            return (f"Câu '{_short(closing[0])}' có dấu hiệu {', '.join(labels)} – "
                    "đây có thật là quyết định đã chốt không?")
        return f"Code và AI hiểu khác nhau về '{content}' – trạng thái đúng là gì?"
    if reason == "possible_action":
        return f"'{content}' là quyết định hay việc giao cho {item.owner or 'ai đó'}?"
    if reason == "possible_closing":
        found = _utt_text(result, _note_utterance(item, ("possible_closing",))) or closing
        quote = _short(found[0]) if found else _first_quote(item, result)
        return f"Câu '{quote}' có phải là câu chốt cho '{content}' không?"
    if reason == "possible_superseded":
        found = _utt_text(result, _note_utterance(item, ("reversal", "sibling")))
        if found:
            return f"Câu '{_short(found[0])}' ở phút {found[1]} có phải đã thay đổi quyết định này không?"
        return f"Có câu nào sau đó đã thay đổi quyết định '{content}' không?"
    if reason == "authority_unconfirmed":
        by_id = {u.id: u.speaker for u in result.utterances}
        who = by_id.get(_closing_uid(item) or "", "Người chốt")
        return f"'{who}' không có quyền chốt. Quyết định này đã được ai xác nhận chưa?"
    if reason == "number_not_in_evidence":
        missing = check.get("missing_numbers")
        number = str(missing[0]) if isinstance(missing, list) and missing else _number_in(item.content)
        if number:
            return f"Số '{number}' không có trong câu gốc. Số đúng là bao nhiêu?"
        return f"Có số trong '{content}' không có trong câu gốc. Số đúng là bao nhiêu?"
    return None


def clarification_question(item: Item, reason: str, result: RunResult) -> str:
    """Vietnamese follow-up question for one reason, from a fixed template (no AI)."""
    content = _short(item.content)
    signal_question = _signal_clarification(item, reason, result, content)
    if signal_question:
        return signal_question
    if reason == "missing_owner":
        return f"Ai chịu trách nhiệm: {content}?"
    if reason == "unresolved_reference":
        return f"'{_first_quote(item, result)}' đang nói tới phương án nào?"
    if reason == "conflicting_decision":
        return f"'{content}' và '{_other_decision_text(item, result)}' – quyết định nào là cuối cùng?"
    if reason == "needs_recheck":
        number = _number_in(item.content)
        if number:
            return f"Con số '{number}' còn đúng không?"
        return f"Thông tin '{content}' còn đúng không?"
    if reason == "quote_mismatch":
        return f"Trích dẫn '{_first_quote(item, result)}' không khớp câu gốc – nội dung '{content}' có đúng không?"
    if reason == NO_FOLLOW_UP:
        return f"Ai theo dõi và báo lại (đang chờ: {item.pending_on or 'chưa rõ'}): {content}?"
    return f"{content} – cần làm rõ điều gì?"


def clarification_reasons(item: Item) -> list[str]:
    reasons = [f for f in CLARIFY_FLAGS if f in item.flags]
    status = item.status.value if item.status else None
    if status in ("deferred", "conditional") and not item.follow_up_owner:
        reasons.append(NO_FOLLOW_UP)
    return reasons


def _answer_field(reason: str) -> str:
    return {"missing_owner": "owner", NO_FOLLOW_UP: "follow_up_owner"}.get(reason, "note")


def clarification_items(result: RunResult, reviews: dict[str, Any] | None = None) -> list[Clarification]:
    """Questions for items needing clarification, in transcript order of items.

    Skipped: rejected items and questions the reviewer marked resolved (kept in reviews.json). `result` should be
    the review-applied result, so a filled owner / follow_up_owner already removes its question.
    """
    reviews = reviews or {}
    out: list[Clarification] = []
    for item in result.items:
        if item.review == "rejected":
            continue
        entry = reviews.get("items", {}).get(item.item_id, {})
        answers = entry.get("answers") if isinstance(entry.get("answers"), dict) else {}
        resolved = entry.get("resolved") if isinstance(entry.get("resolved"), dict) else {}
        for reason in clarification_reasons(item):
            if resolved.get(reason):
                continue
            out.append(Clarification(item.item_id, reason, clarification_question(item, reason, result),
                                     _answer_field(reason), answers.get(reason) or None))
    return out


def resolve_edits(existing: dict[str, Any], reason: str) -> dict[str, Any]:
    """Review-entry fields recording that the question `reason` is resolved (does not touch `review`)."""
    resolved = dict(existing.get("resolved") or {})
    resolved[reason] = True
    return {"resolved": resolved}


def answer_edits(existing: dict[str, Any], c: Clarification, answer: str) -> dict[str, Any]:
    """Fields to merge into a review entry when the reviewer answers `c` (caller sets review='edited')."""
    answer = answer.strip()
    answers = dict(existing.get("answers") or {})
    answers[c.reason] = answer
    edits: dict[str, Any] = {"answers": answers, "note": "\n".join(f"{k}: {v}" for k, v in answers.items())}
    if c.answer_field != "note":
        edits[c.answer_field] = answer
    return edits


# ---------------------------------------------------------------- knowledge base


def knowledge_entries(result: RunResult, only_reviewed: bool = False) -> list[dict[str, Any]]:
    """Knowledge items with scope project/reusable (never ephemeral, never rejected)."""
    meta = result.meta
    meeting = (meta.title if meta and meta.title else None) or Path(result.source_file).name
    by_id = {u.id: u for u in result.utterances}
    out: list[dict[str, Any]] = []
    for item in result.items:
        if item.type != ItemType.knowledge or item.knowledge_scope not in KNOWLEDGE_SCOPES_KEPT:
            continue
        if item.review == "rejected" or (only_reviewed and item.review not in EXPORTED_REVIEWS):
            continue
        utts = [{"id": u.id, "time": u.time, "speaker": u.speaker, "text": u.text}
                for u in (by_id[e] for e in item.evidence if e in by_id)]
        out.append({
            "item_id": item.item_id, "content": item.content, "topic": item.topic,
            "scope": item.knowledge_scope, "scope_label": SCOPE_LABELS[item.knowledge_scope],
            "needs_recheck": "needs_recheck" in item.flags, "confidence": round(item.confidence, 2),
            "review": item.review,
            "source": {"meeting": meeting, "date": meta.date if meta else None,
                       "project": meta.project if meta else None, "utterances": utts},
        })
    return out


def knowledge_to_json(result: RunResult, only_reviewed: bool = False) -> str:
    return json.dumps({"run_id": result.run_id, "entries": knowledge_entries(result, only_reviewed)},
                      ensure_ascii=False, indent=2)


KNOWLEDGE_COLUMNS = [
    ("ID", 7), ("Nội dung", 52), ("Chủ đề", 22), ("Phạm vi", 18), ("Cần kiểm tra lại", 16),
    ("Cuộc họp", 28), ("Ngày", 12), ("Câu nguồn", 60), ("Độ tin cậy", 11), ("Duyệt", 12),
]


def _add_knowledge_sheet(wb: Workbook, result: RunResult, only_reviewed: bool, header_fill: PatternFill) -> None:
    entries = knowledge_entries(result, only_reviewed)
    if not entries:
        return
    ws = wb.create_sheet(KNOWLEDGE_SHEET)
    ws.append([c for c, _ in KNOWLEDGE_COLUMNS])
    for idx, (_, width) in enumerate(KNOWLEDGE_COLUMNS, start=1):
        cell = ws.cell(row=1, column=idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        ws.column_dimensions[get_column_letter(idx)].width = width
    for e in entries:
        src = e["source"]
        cites = "\n".join(f"{u['id']} [{u['time']}] {u['speaker']}: {u['text']}" for u in src["utterances"])
        append_row(ws, [e["item_id"], e["content"], e["topic"], e["scope_label"], "Có" if e["needs_recheck"] else "",
                   src["meeting"], src["date"] or "", cites, e["confidence"],
                   REVIEW_LABELS.get(e["review"], e["review"])])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
