"""End-to-end pipeline with FakeLLM – no API calls."""

import json
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from meeting_extractor.config import Settings
from meeting_extractor.export import apply_reviews, export_run, load_reviews, save_reviews, to_excel
from meeting_extractor.llm import FakeLLM
from meeting_extractor.models import Status
from meeting_extractor.parsing import load_meta
from meeting_extractor.pipeline import load_result, run

FIXTURES = Path(__file__).parent / "fixtures"

T01_IDS = ["u006", "u007", "u008", "u009", "u010", "u011", "u014", "u015", "u016", "u019"]


@pytest.fixture
def settings(tmp_path) -> Settings:
    return Settings(anthropic_api_key=None, model_light="light", model_strong="strong",
                    temperature=None, runs_dir=tmp_path / "runs")


def _extract(system: str, user: str) -> dict:
    if "t01" in user:
        return {"candidates": [
            {"type": "decision", "topic_id": "t01", "content": "On-premise", "speaker": "Nam", "evidence": ["u006"]},
            {"type": "decision", "topic_id": "t01", "content": "AWS", "speaker": "Lan", "evidence": ["u015"]},
        ]}
    return {"candidates": [
        {"type": "knowledge", "topic_id": "t02", "content": "Xong wireframe", "speaker": "Minh", "evidence": ["u004"]},
    ]}


def _track(system: str, user: str) -> dict:
    if "t01" in user:
        return {"items": [{
            "type": "decision", "content": "Dùng AWS cho hạ tầng server", "status": "confirmed",
            "history": [
                {"status": "proposed", "content": "On-premise", "evidence": ["u006"]},
                {"status": "superseded", "content": "On-premise", "reason": "vượt ngân sách quý",
                 "evidence": ["u014"]},
                {"status": "confirmed", "content": "AWS", "evidence": ["u019"]},
            ],
            "evidence": ["u019"], "confidence": 0.92,
        }]}
    return {"items": [
        {"type": "knowledge", "content": "Đã xong wireframe đăng nhập", "evidence": ["u004"], "confidence": 0.9},
    ]}


def tc1_llm() -> FakeLLM:
    return FakeLLM({
        "segment": {"topics": [
            {"topic_id": "t01", "title": "Hạ tầng server", "utterance_ids": T01_IDS},
            {"topic_id": "t02", "title": "Tiến độ thiết kế", "utterance_ids": ["u004", "u012"]},
        ]},
        "extract": _extract,
        "track": _track,
        "verify": {"passed": True, "reason": "Câu u019 chốt AWS"},
    })


def test_pipeline_tc1_end_to_end(settings):
    steps: list[float] = []
    result = run(FIXTURES / "tc1.txt", load_meta(FIXTURES / "tc1.meta.json"), llm=tc1_llm(),
                 settings=settings, progress=lambda msg, frac: steps.append(frac))
    decision = next(i for i in result.items if i.type.value == "decision")
    assert decision.item_id == "d01"
    assert decision.status == Status.confirmed and decision.evidence == ["u019"]
    assert decision.history[1].status == Status.superseded
    assert not any(i.status == Status.confirmed and "premise" in i.content.lower() for i in result.items)
    assert {"segment", "extract", "track", "verify"} <= set(result.prompt_versions)
    assert steps[-1] == 1.0

    run_dir = settings.runs_dir / result.run_id
    for name in ["result.json", "segment.json", "extract_t01.json", "extract_t02.json", "usage.json"]:
        assert (run_dir / name).exists(), name
    assert load_result(run_dir).items == result.items
    usage = json.loads((run_dir / "usage.json").read_text(encoding="utf-8"))
    assert usage["calls"] == 6  # segment + 2×extract + 2×track + 1 verify


def test_pipeline_survives_llm_failure_on_one_topic(settings):
    llm = tc1_llm()
    llm.handlers["track"] = lambda s, u: "không phải json" if "t02" in u else _track(s, u)
    result = run(FIXTURES / "tc1.txt", llm=llm, settings=settings)
    assert any("t02" in w for w in result.warnings)
    # Phase 1: the failed topic keeps Agent 2's raw candidates (flagged) instead of being dropped.
    assert len(result.items) == 2
    assert result.failed_topics == ["t02"]
    assert any("tracker_failed" in i.flags for i in result.items if i.topic_id == "t02")


def test_pipeline_no_meta_still_runs(settings):
    result = run(FIXTURES / "tc1.txt", None, llm=tc1_llm(), settings=settings, save=False)
    assert result.meta is None and result.items
    assert not settings.runs_dir.exists()


def test_reviews_and_excel_export(settings):
    result = run(FIXTURES / "tc1.txt", llm=tc1_llm(), settings=settings)
    run_dir = settings.runs_dir / result.run_id
    reviews = load_reviews(run_dir)
    reviews["items"]["d01"] = {"review": "edited", "owner": "Speaker 9", "content": "Dùng AWS (trả theo tháng)"}
    reviews["items"]["k01"] = {"review": "rejected"}
    reviews["speaker_names"] = {"Speaker 9": "Tuấn"}
    save_reviews(run_dir, reviews)

    reviewed = apply_reviews(load_result(run_dir), load_reviews(run_dir))
    d01 = next(i for i in reviewed.items if i.item_id == "d01")
    assert (d01.review, d01.content, d01.owner) == ("edited", "Dùng AWS (trả theo tháng)", "Tuấn")

    wb = load_workbook(BytesIO(to_excel(reviewed)))
    assert "Quyết định" in wb.sheetnames and "Tri thức mới" not in wb.sheetnames
    assert export_run(run_dir, "xlsx").exists()
    assert export_run(run_dir, "json").exists()


def test_cli_run_and_export(settings, monkeypatch, capsys):
    import meeting_extractor.cli as cli
    import meeting_extractor.pipeline as pipeline

    monkeypatch.setattr(pipeline, "load_settings", lambda: settings)
    monkeypatch.setattr(pipeline, "make_llm", lambda s: tc1_llm())
    assert cli.main(["run", str(FIXTURES / "tc1.txt"), "--meta", str(FIXTURES / "tc1.meta.json")]) == 0
    out = capsys.readouterr().out
    assert "d01" in out and "confirmed" in out
    run_dir = next(settings.runs_dir.iterdir())
    assert cli.main(["export", str(run_dir), "--format", "xlsx", "--all"]) == 0
