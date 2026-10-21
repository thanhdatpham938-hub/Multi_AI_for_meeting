"""Mission 3 / 5B: evidence display, clarification tab, knowledge base, related items, old runs (offline)."""

import html
import json
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from meeting_extractor.export import (
    NO_FOLLOW_UP,
    answer_edits,
    apply_reviews,
    clarification_items,
    empty_reviews,
    export_run,
    highlight_quote_html,
    knowledge_entries,
    knowledge_to_json,
    load_reviews,
    locate_quote,
    related_summaries,
    to_excel,
    utterance_context,
)
from meeting_extractor.models import Item, MeetingMeta, RunResult, Utterance
from tests.test_phase4_ui import _make_run, _open_detail, _open_run, _patch_result, _text, _audio_uploader, APP

@pytest.fixture
def run_dir(tmp_path, monkeypatch) -> Path:
    return _make_run(tmp_path, monkeypatch)


# ---------------------------------------------------------------- builders


def _utts(n: int = 7) -> list[Utterance]:
    return [Utterance(id=f"u{k:03d}", time=f"00:{k:02d}", speaker="Hùng" if k % 2 else "Lan", text=f"Câu số {k}.")
            for k in range(1, n + 1)]


def _item(item_id: str, type_: str = "decision", **kw) -> Item:
    base = dict(item_id=item_id, type=type_, topic_id="t01", topic="Chủ đề", content=f"Nội dung {item_id}",
                evidence=["u004"], confidence=0.8)
    base.update(kw)
    return Item(**base)


def _result(items: list[Item], utts: list[Utterance] | None = None, meta: MeetingMeta | None = None) -> RunResult:
    return RunResult(run_id="r1", source_file="x.txt", model_light="l", model_strong="s", prompt_versions={},
                     utterances=utts or _utts(), topics=[], items=items, meta=meta)


# ---------------------------------------------------------------- highlight


def test_highlight_bolds_quote_after_normalization():
    html_, found = highlight_quote_html("OK, vậy mình   chốt dùng AWS nhé.", "Chốt dùng aws")
    assert found is True
    assert html_ == "OK, vậy mình   <b>chốt dùng AWS</b> nhé."


def test_highlight_quote_with_edge_punctuation_and_unicode_forms():
    import unicodedata

    assert highlight_quote_html("Vậy mình chốt dùng AWS.", "…chốt dùng AWS.")[1] is True
    decomposed = unicodedata.normalize("NFD", "chốt")  # decomposed diacritics must still match
    assert highlight_quote_html("Vậy mình chốt", decomposed) == ("Vậy mình <b>chốt</b>", True)


def test_highlight_is_html_safe():
    html_, found = highlight_quote_html("<script>alert(1)</script> chốt <b>AWS</b> & xong", "AWS")
    assert found is True
    assert "<script>" not in html_ and "&lt;script&gt;" in html_
    assert "&lt;b&gt;<b>AWS</b>&lt;/b&gt;" in html_ and "&amp; xong" in html_


def test_highlight_not_found_and_no_quote():
    assert highlight_quote_html("Câu a <i>", "không có") == ("Câu a &lt;i&gt;", False)
    assert highlight_quote_html("Câu a", None) == ("Câu a", None)
    assert highlight_quote_html("Câu a", "   ") == ("Câu a", None)
    assert locate_quote("abc", "") is None


def test_utterance_context_edges():
    utts = _utts()
    before, after = utterance_context(utts, "u004")
    assert [u.id for u in before] == ["u002", "u003"] and [u.id for u in after] == ["u005", "u006"]
    before, after = utterance_context(utts, "u001")
    assert before == [] and [u.id for u in after] == ["u002", "u003"]
    assert [u.id for u in utterance_context(utts, "u007")[1]] == []
    assert utterance_context(utts, "zzz") == ([], [])


# ---------------------------------------------------------------- dialog


def _decision_patch(d: dict) -> None:
    ids = [u["id"] for u in d["utterances"]]
    item = d["items"][0]
    quote_id = item["evidence"][0]
    text = next(u["text"] for u in d["utterances"] if u["id"] == quote_id)
    item["evidence_quotes"] = {quote_id: text.split()[0].strip(".,!?:;")}
    item["field_evidence"] = {"owner": [ids[0]], "contradicted_by": [ids[1]]}
    item["owner"] = "Minh"
    item["commitment"] = "assigned"
    item["pending_on"] = "phản hồi của khách"
    item["follow_up_owner"] = "Lan"
    item["related_items"] = [d["items"][1]["item_id"], "x99"]


def test_dialog_shows_bold_quote_context_sources_and_related(run_dir):
    _patch_result(run_dir, _decision_patch)
    data = json.loads((run_dir / "result.json").read_text(encoding="utf-8"))
    ev = data["items"][0]["evidence"][0]
    first_word = data["items"][0]["evidence_quotes"][ev]  # punctuation already stripped
    at = _open_detail(_open_run(run_dir))
    assert not at.exception
    md = _text(at)
    assert f"<b>{first_word}</b>" in md
    assert "quote ctx" in md  # neighbouring sentences, dimmed
    assert "Phụ trách</b> ←" in md and "Bị phản bác bởi" in md
    assert "Mục liên quan" in md and data["items"][1]["item_id"] in md and "(mục không còn trong kết quả)" in md
    assert "Được giao việc" in md and "phản hồi của khách" in md and "Người theo dõi" in md


def test_dialog_context_is_two_before_and_two_after(run_dir):
    at = _open_detail(_open_run(run_dir))
    data = json.loads((run_dir / "result.json").read_text(encoding="utf-8"))
    ev = data["items"][0]["evidence"][0]
    pos = [u["id"] for u in data["utterances"]].index(ev)
    expected = min(pos, 2) + min(len(data["utterances"]) - pos - 1, 2)
    quote_blocks = [m.value for m in at.markdown if "class='quote ctx'" in m.value or 'class="quote ctx"' in m.value]
    assert sum(v.count("quote ctx") for v in quote_blocks) >= expected > 0


def test_dialog_quote_not_found_shows_small_warning_and_flag_warning(run_dir):
    def patch(d):
        item = d["items"][0]
        item["evidence_quotes"] = {item["evidence"][0]: "đoạn này không hề có trong câu"}
        item["flags"] = ["quote_mismatch"]
    _patch_result(run_dir, patch)
    at = _open_detail(_open_run(run_dir))
    assert not at.exception
    assert "Không tìm thấy đoạn trích" in _text(at)
    assert any("không khớp nguyên văn" in w.value for w in at.warning)


def test_dialog_escapes_html_in_content_and_utterance(run_dir):
    def patch(d):
        item = d["items"][0]
        item["content"] = "<img src=x onerror=alert(1)>"
        for u in d["utterances"]:
            if u["id"] == item["evidence"][0]:
                u["text"] = "<script>boom()</script> chốt"
        item["evidence_quotes"] = {item["evidence"][0]: "chốt"}
    _patch_result(run_dir, patch)
    at = _open_detail(_open_run(run_dir))
    md = _text(at)
    assert not at.exception and "<script>" not in md and "<img" not in md and "&lt;script&gt;" in md


def test_old_run_without_new_fields_has_no_error(run_dir):
    def strip(d):
        for item in d["items"]:
            for k in ("commitment", "evidence_quotes", "field_evidence", "knowledge_scope", "pending_on",
                      "follow_up_owner", "related_items"):
                item.pop(k, None)
    _patch_result(run_dir, strip)
    at = _open_detail(_open_run(run_dir))
    assert not at.exception
    md = _text(at)
    assert "Mục liên quan" not in md and "Bị phản bác bởi" not in md and "Không tìm thấy đoạn trích" not in md
    assert "Cần làm rõ" in [t.label for t in at.tabs]


# ---------------------------------------------------------------- clarification (pure)


def test_clarification_lists_only_relevant_items_with_template_questions():
    utts = _utts()
    items = [
        _item("a01", "action", content="Làm tài liệu API", flags=["missing_owner"]),
        _item("d01", content="Dùng AWS", flags=["conflicting_decision"], related_items=["d02"]),
        _item("d02", content="Dùng on-premise"),
        _item("d03", content="Phương án đó", flags=["unresolved_reference"], evidence_quotes={"u004": "phương án đó"}),
        _item("d04", content="Ngân sách 30 triệu", flags=["needs_recheck"]),
        _item("d05", content="Chốt X", flags=["quote_mismatch"], evidence_quotes={"u004": "câu sai"}),
        _item("d06", content="Hoãn ngân sách", status="deferred", pending_on="sếp"),
        _item("d07", content="Hoãn có người", status="deferred", follow_up_owner="Lan"),
        _item("d08", content="Bình thường", status="confirmed"),
        _item("d09", content="Điều kiện", status="conditional", condition="nếu khách đồng ý"),
        _item("d10", content="Đã loại", flags=["missing_owner"], review="rejected"),
    ]
    got = {(c.item_id, c.reason): c for c in clarification_items(_result(items, utts))}
    assert set(got) == {("a01", "missing_owner"), ("d01", "conflicting_decision"), ("d03", "unresolved_reference"),
                        ("d04", "needs_recheck"), ("d05", "quote_mismatch"), ("d06", NO_FOLLOW_UP),
                        ("d09", NO_FOLLOW_UP)}
    assert got[("a01", "missing_owner")].question == "Ai chịu trách nhiệm: Làm tài liệu API?"
    assert got[("a01", "missing_owner")].answer_field == "owner"
    assert got[("d03", "unresolved_reference")].question == "'phương án đó' đang nói tới phương án nào?"
    assert got[("d01", "conflicting_decision")].question == "Dùng AWS và Dùng on-premise – quyết định nào là cuối cùng?"
    assert got[("d04", "needs_recheck")].question == "Con số '30 triệu' còn đúng không?"
    assert "câu sai" in got[("d05", "quote_mismatch")].question
    assert "đang chờ: sếp" in got[("d06", NO_FOLLOW_UP)].question
    assert got[("d06", NO_FOLLOW_UP)].answer_field == "follow_up_owner"


def test_clarification_conflict_without_related_and_text_recheck():
    items = [_item("d01", content="Dùng AWS", flags=["conflicting_decision"]),
             _item("k01", "knowledge", content="Sandbox hết hạn sớm", flags=["needs_recheck"])]
    qs = {c.item_id: c.question for c in clarification_items(_result(items))}
    assert "một quyết định khác" in qs["d01"] and "quyết định nào là cuối cùng" in qs["d01"]
    assert qs["k01"] == "Thông tin 'Sandbox hết hạn sớm' còn đúng không?"


def test_answer_edits_owner_and_note_and_accumulation():
    items = [_item("d01", flags=["needs_recheck", "missing_owner"])]
    by_reason = {c.reason: c for c in clarification_items(_result(items))}
    c_owner, c_note = by_reason["missing_owner"], by_reason["needs_recheck"]
    e1 = answer_edits({}, c_owner, " Minh ")
    assert e1["owner"] == "Minh" and e1["answers"] == {"missing_owner": "Minh"}
    e2 = answer_edits({"answers": e1["answers"]}, c_note, "Vẫn đúng")
    assert "owner" not in e2 and e2["answers"] == {"missing_owner": "Minh", "needs_recheck": "Vẫn đúng"}
    assert "needs_recheck: Vẫn đúng" in e2["note"]


def test_apply_reviews_tolerates_note_and_answers_keys():
    base = _result([_item("a01", "action", flags=["missing_owner"])])
    reviews = {"items": {"a01": {"review": "edited", "note": "ghi chú", "answers": {"x": "y"}, "owner": "Minh",
                                 "unknown_key": 1}}, "speaker_names": {}}
    out = apply_reviews(base, reviews)
    assert out.items[0].owner == "Minh" and out.items[0].review == "edited"


def test_answered_question_is_marked_answered():
    items = [_item("a01", "action", flags=["missing_owner"])]
    reviews = empty_reviews()
    reviews["items"]["a01"] = {"review": "edited", "answers": {"missing_owner": "Minh"}}
    (c,) = clarification_items(_result(items), reviews)
    assert c.answered == "Minh"


def test_answered_follow_up_question_stays_visible_after_follow_up_owner_filled():
    items = [_item("d06", status="deferred", follow_up_owner="Lan")]
    reviews = {"items": {"d06": {"answers": {NO_FOLLOW_UP: "Lan"}}}, "speaker_names": {}}
    (c,) = clarification_items(_result(items), reviews)
    assert c.reason == NO_FOLLOW_UP and c.answered == "Lan"


# ---------------------------------------------------------------- clarification tab (AppTest)


def _flag_first_action_without_owner(run_dir: Path) -> str:
    def patch(d):
        item = d["items"][0]
        item["type"] = "action"
        item["status"] = None
        item["owner"] = None
        item["flags"] = ["missing_owner"]
    _patch_result(run_dir, patch)
    return json.loads((run_dir / "result.json").read_text(encoding="utf-8"))["items"][0]["item_id"]


def test_clarify_tab_lists_item_and_answer_is_saved_as_owner(run_dir):
    item_id = _flag_first_action_without_owner(run_dir)
    at = _open_run(run_dir)
    assert not at.exception
    assert "Ai chịu trách nhiệm:" in _text(at)
    key = f"{item_id}-missing_owner"
    next(t for t in at.text_input if t.key == f"clarify-answer-{key}").set_value("Minh")
    next(b for b in at.button if b.key == f"clarify-submit-{key}").click()
    at.run()
    assert not at.exception
    saved = load_reviews(run_dir)["items"][item_id]
    assert saved["owner"] == "Minh" and saved["review"] == "edited"
    assert saved["answers"]["missing_owner"] == "Minh" and "missing_owner: Minh" in saved["note"]
    assert "Đã trả lời: Minh" in _text(at)
    reviews_file = json.loads((run_dir / "reviews.json").read_text(encoding="utf-8"))
    assert reviews_file["items"][item_id]["owner"] == "Minh"


def test_clarify_tab_non_owner_answer_goes_to_note_only(run_dir):
    def patch(d):
        d["items"][0]["flags"] = ["needs_recheck"]
        d["items"][0]["content"] = "Ngân sách 30 triệu"
    _patch_result(run_dir, patch)
    item_id = json.loads((run_dir / "result.json").read_text(encoding="utf-8"))["items"][0]["item_id"]
    at = _open_run(run_dir)
    assert "Con số '30 triệu' còn đúng không?" in html.unescape(_text(at))  # shown HTML-escaped
    key = f"{item_id}-needs_recheck"
    next(t for t in at.text_input if t.key == f"clarify-answer-{key}").set_value("Đúng, 30 triệu")
    next(b for b in at.button if b.key == f"clarify-submit-{key}").click()
    at.run()
    saved = load_reviews(run_dir)["items"][item_id]
    assert "owner" not in saved and "Đúng, 30 triệu" in saved["note"]
    assert not at.exception


def test_clarify_tab_blank_answer_not_saved(run_dir):
    item_id = _flag_first_action_without_owner(run_dir)
    at = _open_run(run_dir)
    next(b for b in at.button if b.key == f"clarify-submit-{item_id}-missing_owner").click()
    at.run()
    assert not at.exception and item_id not in load_reviews(run_dir)["items"]


# ---------------------------------------------------------------- knowledge base


def _knowledge_result() -> RunResult:
    meta = MeetingMeta(title="Họp dự án X", date="12/04/2025", project="X")
    items = [
        _item("k01", "knowledge", content="Sandbox hết hạn sau 15 phút", knowledge_scope="reusable",
              evidence=["u002", "u003"], review="accepted"),
        _item("k02", "knowledge", content="Staging hôm qua chậm", knowledge_scope="ephemeral", review="accepted"),
        _item("k03", "knowledge", content="API giới hạn 100 req/s", knowledge_scope="project",
              flags=["needs_recheck"], review="edited"),
        _item("k04", "knowledge", content="Chưa gắn nhãn", review="accepted"),
        _item("k05", "knowledge", content="Đã loại", knowledge_scope="project", review="rejected"),
        _item("k06", "knowledge", content="Chưa duyệt", knowledge_scope="project"),
        _item("d01", content="Không phải tri thức", knowledge_scope="project", review="accepted"),
    ]
    return _result(items, meta=meta)


def test_knowledge_entries_only_project_and_reusable():
    entries = knowledge_entries(_knowledge_result())
    assert [e["item_id"] for e in entries] == ["k01", "k03", "k06"]  # no ephemeral, unlabeled, rejected, non-knowledge
    k01 = entries[0]
    assert k01["scope"] == "reusable" and k01["needs_recheck"] is False
    assert k01["source"]["meeting"] == "Họp dự án X" and k01["source"]["date"] == "12/04/2025"
    assert [u["id"] for u in k01["source"]["utterances"]] == ["u002", "u003"]
    assert k01["source"]["utterances"][0]["speaker"] and k01["source"]["utterances"][0]["time"]
    assert entries[1]["needs_recheck"] is True
    assert [e["item_id"] for e in knowledge_entries(_knowledge_result(), only_reviewed=True)] == ["k01", "k03"]


def test_knowledge_json_content():
    data = json.loads(knowledge_to_json(_knowledge_result(), only_reviewed=True))
    assert data["run_id"] == "r1" and [e["item_id"] for e in data["entries"]] == ["k01", "k03"]
    assert data["entries"][1]["scope_label"] == "Theo dự án"


def test_excel_has_knowledge_sheet_with_columns():
    wb = load_workbook(BytesIO(to_excel(_knowledge_result())))
    assert "Kho tri thức" in wb.sheetnames
    ws = wb["Kho tri thức"]
    header = [c.value for c in ws[1]]
    assert header[:5] == ["ID", "Nội dung", "Chủ đề", "Phạm vi", "Cần kiểm tra lại"] and "Câu nguồn" in header
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    assert [r[0] for r in rows] == ["k01", "k03"]
    assert rows[0][3] == "Dùng lại nhiều lần" and rows[1][4] == "Có"
    assert "Họp dự án X" in rows[0] and "12/04/2025" in rows[0] and "u002" in rows[0][7]


def test_excel_without_knowledge_has_no_sheet():
    res = _result([_item("d01", review="accepted")])
    assert "Kho tri thức" not in load_workbook(BytesIO(to_excel(res))).sheetnames


def test_export_run_knowledge_writes_json(run_dir):
    def patch(d):
        d["items"][0].update(type="knowledge", status=None, knowledge_scope="project")
    _patch_result(run_dir, patch)
    path = export_run(run_dir, "knowledge", only_reviewed=False)
    assert path.name == "knowledge.json"
    assert len(json.loads(path.read_text(encoding="utf-8"))["entries"]) == 1


def test_knowledge_tab_shows_entries_and_download(run_dir):
    def patch(d):
        d["items"][0].update(type="knowledge", status=None, knowledge_scope="reusable",
                             content="Sandbox hết hạn sau 15 phút")
    _patch_result(run_dir, patch)
    at = _open_run(run_dir)
    assert not at.exception
    assert "Kho tri thức" in [t.label for t in at.tabs]
    assert any("Sandbox hết hạn sau 15 phút" in str(df.value.to_dict()) for df in at.dataframe)


def test_knowledge_tab_old_run_shows_placeholder_not_error(run_dir):
    def patch(d):
        d["items"][0].update(type="knowledge", status=None)
        d["items"][0].pop("knowledge_scope", None)
    _patch_result(run_dir, patch)
    at = _open_run(run_dir)
    assert not at.exception
    assert "chưa có nhãn phạm vi" in _text(at)
    assert any("Chưa có tri thức nào" in i.value for i in at.info)


# ---------------------------------------------------------------- related items


def test_related_summaries_truncate_and_missing():
    long = "x" * 200
    res = _result([_item("d01", related_items=["a01", "zz"]), _item("a01", "action", content=long)])
    out = related_summaries(res.items[0], res, max_len=20)
    assert out[0][0] == "a01" and out[0][2] is True and len(out[0][1]) == 20 and out[0][1].endswith("…")
    assert out[1] == ("zz", "", False)


# ---------------------------------------------------------------- N1 / N2 / N3 / N6


def test_n1_closer_unknown_label(run_dir):
    _patch_result(run_dir, lambda d: d["items"][0].update(confidence_breakdown={"base": 0.5, "closer_unknown": -0.1}))
    at = _open_detail(_open_run(run_dir))
    assert "người chốt không có quyền hoặc chưa xác định" in _text(at)


def test_n2_confidence_uses_decimal_comma_in_table_dialog_and_help(run_dir):
    _patch_result(run_dir, lambda d: d["items"][0].update(confidence=0.72))
    at = _open_run(run_dir)
    assert "tin cậy 0,72" in _text(at) and "72%" not in _text(at)
    assert any("≥ 0,70" in (b.help or "") for b in at.button)
    assert "0,72" in _text(_open_detail(at))


def test_n3_oversize_audio_rejected_without_writing_to_disk(monkeypatch, tmp_path):
    monkeypatch.setattr("meeting_extractor.audio.MAX_AUDIO_BYTES", 50)
    writes: list = []
    real = Path.write_bytes
    monkeypatch.setattr(Path, "write_bytes", lambda self, data: writes.append(self) or real(self, data))
    at = AppTest.from_file(APP, default_timeout=30).run()
    _audio_uploader(at).upload("big.wav", b"0" * 200, "audio/wav")
    at.run()
    at.run()  # a rerun must not write either
    assert not at.exception and any("quá lớn" in e.value for e in at.error)
    assert not any("meetsync_" in str(p) for p in writes)


def test_n3_valid_audio_not_written_until_analyze_clicked(monkeypatch):
    writes: list = []
    real = Path.write_bytes
    monkeypatch.setattr(Path, "write_bytes", lambda self, data: writes.append(self) or real(self, data))
    at = AppTest.from_file(APP, default_timeout=30).run()
    _audio_uploader(at).upload("a.mp3", b"ID3" + b"0" * 10, "audio/mpeg")
    at.run()
    at.run()
    assert not at.exception and not any("meetsync_" in str(p) for p in writes)


def test_n6_oserror_when_saving_audio_shows_vietnamese_error(monkeypatch):
    def boom(self, data):
        raise OSError("disk full")

    at = AppTest.from_file(APP, default_timeout=30).run()
    _audio_uploader(at).upload("a.mp3", b"ID3" + b"0" * 10, "audio/mpeg")
    at.run()
    monkeypatch.setattr(Path, "write_bytes", boom)
    next(b for b in at.button if "Phân tích" in b.label).click()
    at.run()
    assert not at.exception
    assert any("Không ghi được file" in e.value and "disk full" in e.value for e in at.error)
