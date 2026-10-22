"""Phase S3: signal-layer flags in the UI: clarification questions, two-level flags, code-check block (offline)."""

import html
import json
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from meeting_extractor.export import (
    HIGH_LEVEL_FLAGS,
    apply_reviews,
    clarification_items,
    empty_reviews,
    md_safe,
    split_flags,
    to_excel,
)
from meeting_extractor.models import Item, RunResult, Utterance
from tests.test_phase4_ui import _make_run, _open_detail, _open_run, _patch_result, _text


@pytest.fixture
def run_dir(tmp_path, monkeypatch) -> Path:
    return _make_run(tmp_path, monkeypatch)


UTTS = [
    Utterance(id="u001", time="10:00", speaker="Lan", text="Mình dùng AWS nhé"),
    Utterance(id="u002", time="10:05", speaker="Nam", text="Đùa thôi, chưa chốt đâu"),
    Utterance(id="u003", time="12:30", speaker="Hùng", text="Thôi đổi sang GCP, 200 triệu"),
]


def _item(flags: list[str], signal_check: dict | None = None, **kw) -> Item:
    base = dict(item_id="d01", type="decision", topic_id="t01", topic="Hạ tầng", content="Dùng AWS",
                status="confirmed", evidence=["u001"], confidence=0.5, flags=flags,
                signal_check=signal_check or {})
    base.update(kw)
    return Item(**base)


def _question(item: Item) -> str:
    result = RunResult(run_id="r", source_file="x.txt", model_light="l", model_strong="s", prompt_versions={},
                       utterances=UTTS, topics=[], items=[item])
    (c,) = clarification_items(result)
    assert c.reason == item.flags[0] and c.answer_field == "note"
    return c.question


def _item_id(run_dir: Path) -> str:
    return json.loads((run_dir / "result.json").read_text(encoding="utf-8"))["items"][0]["item_id"]


# ---------------------------------------------------------------- questions


def test_signal_mismatch_question_with_and_without_data():
    check = {"closing_utterance": "u002", "tags": {"DUA": "đùa thôi"}, "notes": []}
    assert _question(_item(["signal_mismatch"], check)) == (
        "Câu 'Đùa thôi, chưa chốt đâu' có dấu hiệu đùa – đây có thật là quyết định đã chốt không?")
    assert _question(_item(["signal_mismatch"])) == "Code và AI hiểu khác nhau về 'Dùng AWS' – trạng thái đúng là gì?"
    assert "Code và AI" in _question(_item(["signal_mismatch"], {"tags": "bad", "notes": "bad"}))


def test_possible_action_question():
    assert _question(_item(["possible_action"], owner="Minh")) == "'Dùng AWS' là quyết định hay việc giao cho Minh?"
    assert "giao cho ai đó" in _question(_item(["possible_action"]))


def test_possible_closing_question():
    check = {"notes": [{"rule": "possible_closing", "utterance": "u003", "text": "x"}]}
    assert _question(_item(["possible_closing"], check)) == (
        "Câu 'Thôi đổi sang GCP, 200 triệu' có phải là câu chốt cho 'Dùng AWS' không?")
    assert "Câu 'Mình dùng AWS nhé' có phải" in _question(_item(["possible_closing"]))  # falls back to evidence


def test_possible_superseded_question():
    check = {"notes": [{"rule": "reversal", "utterance": "u003", "text": "x"}]}
    assert _question(_item(["possible_superseded"], check)) == (
        "Câu 'Thôi đổi sang GCP, 200 triệu' ở phút 12:30 có phải đã thay đổi quyết định này không?")
    assert "Có câu nào sau đó đã thay đổi" in _question(_item(["possible_superseded"]))


def test_authority_unconfirmed_question():
    assert _question(_item(["authority_unconfirmed"], {"closing_utterance": "u003"})) == (
        "'Hùng' không có quyền chốt. Quyết định này đã được ai xác nhận chưa?")
    assert "'Lan' không có quyền chốt" in _question(_item(["authority_unconfirmed"]))  # last evidence


def test_number_not_in_evidence_question():
    assert _question(_item(["number_not_in_evidence"], {"missing_numbers": ["200"]})) == (
        "Số '200' không có trong câu gốc. Số đúng là bao nhiêu?")
    assert "Số '30 triệu'" in _question(_item(["number_not_in_evidence"], content="Ngân sách 30 triệu"))
    assert "Có số trong 'Dùng AWS'" in _question(_item(["number_not_in_evidence"]))


def test_split_flags_and_constant():
    high, low = split_flags(["low_confidence", "signal_mismatch", "missing_owner", "verify_failed"])
    assert high == ["signal_mismatch", "verify_failed"] and low == ["low_confidence", "missing_owner"]
    assert "possible_closing" not in HIGH_LEVEL_FLAGS


# ---------------------------------------------------------------- UI


def test_low_flags_collapsed_in_table_and_listed_in_expander(run_dir):
    _patch_result(run_dir, lambda d: d["items"][0].update(flags=["number_check", "low_confidence", "signal_mismatch"]))
    at = _open_run(run_dir)
    assert not at.exception
    assert "+2 ghi chú" in _text(at) and "Lệch tín hiệu code/AI" in _text(at)
    at = _open_detail(at)
    assert not at.exception
    assert any(e.label == "Ghi chú khác (2)" for e in at.expander)
    detail = _text(at)
    assert "Kiểm tra con số" in detail and "Độ tin cậy thấp" in detail


def test_table_hides_low_flag_labels(run_dir):
    _patch_result(run_dir, lambda d: d["items"][0].update(flags=["number_check"]))
    at = _open_run(run_dir)
    assert not at.exception
    assert "+1 ghi chú" in _text(at)
    assert not any("Kiểm tra con số" in m.value and "badge" in m.value for m in at.markdown)


def test_code_check_block_shows_notes_and_utterance_with_md_safe(run_dir):
    uid = json.loads((run_dir / "result.json").read_text(encoding="utf-8"))["utterances"][0]["id"]

    def mutate(d):
        d["items"][0]["signal_check"] = {"notes": [
            {"rule": "weak", "utterance": uid, "text": "Câu có *đùa* và $5 và [x]"},
            {"rule": "number", "utterance": None, "text": "Không câu"},
            {"rule": "x", "text": ""},
        ]}
    _patch_result(run_dir, mutate)
    at = _open_detail(_open_run(run_dir))
    assert not at.exception
    md = _text(at)
    assert "Kiểm tra bằng code" in md and "Không câu" in md
    assert md_safe("Câu có *đùa* và $5 và [x]") in md and "*đùa*" not in md
    assert f"{uid} ·" in md


def test_no_code_check_block_without_notes(run_dir):
    _patch_result(run_dir, lambda d: d["items"][0].update(signal_check={}))
    at = _open_detail(_open_run(run_dir))
    assert not at.exception and "Kiểm tra bằng code" not in _text(at)


def test_new_breakdown_keys_have_vietnamese_labels(run_dir):
    _patch_result(run_dir, lambda d: d["items"][0].update(confidence_breakdown={
        "base": 0.5, "signal_strong_conflict": -0.25, "cap_signal_unverified": -0.05}))
    md = _text(_open_detail(_open_run(run_dir)))
    assert "lệch tín hiệu code (mạnh)" in md and "trần: code không thấy từ chốt" in md


def test_signal_warning_shown_at_top(run_dir):
    _patch_result(run_dir, lambda d: d["warnings"].extend(["[Cảnh báo cờ] 40% decision lệch *tín hiệu*", "khác"]))
    at = _open_run(run_dir)
    assert not at.exception
    assert [w.value for w in at.warning] == [md_safe("[Cảnh báo cờ] 40% decision lệch *tín hiệu*")]


def test_clarify_tab_answers_signal_question_as_note(run_dir):
    _patch_result(run_dir, lambda d: d["items"][0].update(
        flags=["signal_mismatch"], signal_check={"closing_utterance": None, "tags": {}, "notes": []}))
    at = _open_run(run_dir)
    assert "Code và AI hiểu khác nhau" in html.unescape(_text(at))
    key = f"{_item_id(run_dir)}-signal_mismatch"
    next(t for t in at.text_input if t.key == f"clarify-answer-{key}").set_value("Đã chốt thật")
    next(b for b in at.button if b.key == f"clarify-submit-{key}").click()
    at.run()
    assert not at.exception and "Đã trả lời: Đã chốt thật" in html.unescape(_text(at))


def test_excel_keeps_all_flags(run_dir):
    _patch_result(run_dir, lambda d: d["items"][0].update(
        flags=["number_check", "low_confidence", "signal_mismatch"], review="accepted"))
    raw = RunResult.model_validate_json((run_dir / "result.json").read_text(encoding="utf-8"))
    wb = load_workbook(BytesIO(to_excel(apply_reviews(raw, empty_reviews()))))
    cells = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value)
    assert "Kiểm tra con số" in cells and "Độ tin cậy thấp" in cells and "Lệch tín hiệu code/AI" in cells
    assert "ghi chú" not in cells
